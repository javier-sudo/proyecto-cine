import json
import re
from pathlib import Path

from openpyxl import load_workbook


ROOT = Path(__file__).resolve().parents[1]
SOURCE = ROOT / ".source"
OUTPUT = ROOT / "dist" / "catalog-data.js"

SALES = [
    ("camera.xlsx", "Cámara", "cam"),
    ("light.xlsx", "Iluminación", "ilu"),
    ("sound.xlsx", "Sonido", "son"),
    ("production.xlsx", "Producción", "pro"),
]

RENTAL_SECTIONS = {
    "HMI": "HMI",
    "TUNGTENO": "Tungsteno",
    "LED": "LED",
    "TUBOS": "Tubos",
    "OTROS": "Otros",
    "ACCESORIOS": "Accesorios",
}


def clean(value):
    if value is None:
        return ""
    text = str(value).replace("�", "°")
    return re.sub(r"\s+", " ", text).strip()


def sales_category(filename, label, prefix):
    sheet = load_workbook(SOURCE / filename, data_only=True).active
    products = []
    for row in range(1, sheet.max_row + 1):
        name = clean(sheet.cell(row, 2).value)
        if not name:
            continue
        lowered = name.lower()
        if lowered == "kimun spa" or lowered.startswith(("insumos", "total", "iva", "contacto")):
            continue
        products.append({
            "id": f"{prefix}-{len(products) + 1:03d}",
            "name": name,
            "brand": clean(sheet.cell(row, 4).value),
            "detail": clean(sheet.cell(row, 5).value),
            "mode": "Venta",
            "category": label,
        })
    return {"name": label, "products": products}


def rental_categories():
    sheet = load_workbook(SOURCE / "rental.xlsx", data_only=True).active
    grouped = {name: [] for name in RENTAL_SECTIONS.values()}
    section = "Otros"
    for row in range(1, sheet.max_row + 1):
        name = clean(sheet.cell(row, 2).value)
        if not name:
            continue
        upper = name.upper()
        if upper in RENTAL_SECTIONS:
            section = RENTAL_SECTIONS[upper]
            continue
        lowered = name.lower()
        if lowered.startswith(("total", "iva", "contacto", "kimuntecnica")):
            continue
        current = grouped[section]
        current.append({
            "id": f"ren-{section[:3].lower()}-{len(current) + 1:03d}",
            "name": name,
            "brand": clean(sheet.cell(row, 4).value),
            "detail": "",
            "availability": clean(sheet.cell(row, 6).value),
            "mode": "Rental",
            "category": section,
        })
    return [{"name": name, "products": products} for name, products in grouped.items() if products]


catalog = {
    "sale": {"label": "Venta", "categories": [sales_category(*entry) for entry in SALES]},
    "rental": {"label": "Rental", "categories": rental_categories()},
}

OUTPUT.write_text(
    "window.KIMUN_CATALOG = " + json.dumps(catalog, ensure_ascii=False, indent=2) + ";\n",
    encoding="utf-8",
)

print(json.dumps({
    "sale": sum(len(category["products"]) for category in catalog["sale"]["categories"]),
    "rental": sum(len(category["products"]) for category in catalog["rental"]["categories"]),
}, ensure_ascii=False))
